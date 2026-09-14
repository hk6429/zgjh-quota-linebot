# -*- coding: utf-8 -*-
"""
由 knowledge_base.json 產出竹光國中 115學年度總量管制 Google 試算表 (Excel .xlsx)
7 大工作表：Q&A、順位矩陣、學區、日程、歷年門檻、提問紀錄簿、待補充問題庫
用法：python3 create_sheets_db.py  （同時輸出到雲端硬碟資料夾與專案目錄）
"""

import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
KB_PATH = os.path.join(BASE_DIR, 'knowledge_base.json')
FILE_NAME = '竹光國中115學年度總量管制知識庫.xlsx'
GDRIVE_DIR = os.path.expanduser(
    '~/Library/CloudStorage/GoogleDrive-doc614@mail.zgjh.hc.edu.tw/我的雲端硬碟/116 資訊組/03 領域/總量管制相關資料')

QA_HEADERS = ['編號', '主題分類', '觸發關鍵字 (逗號分隔；a+b=同時出現；-詞=排除)', '家長常見諮詢問題',
              '小幫手標準回覆內容 (支援LINE格式)', '法規與依據備註', '啟用狀態']
INQUIRY_HEADERS = ['提問時間', '諮詢單號', '家長LINE識別碼', '學生畢業國小', '家長提問內容', '小幫手AI回覆摘要', '處理狀態', '承辦同仁備註']
UNANSWERED_HEADERS = ['收集時間', '家長原句提問', '觸發模式', '提問者辨識碼', '出現次數', '建議增修處室', '註冊組擬定官方答案', '納入知識庫狀態']


def build_quota_database(kb, output_path):
    wb = Workbook()
    header_fill = PatternFill(start_color='047857', end_color='047857', fill_type='solid')
    header_font = Font(name='微軟正黑體', size=11, bold=True, color='FFFFFF')
    regular_font = Font(name='微軟正黑體', size=10, color='1F2937')
    thin = Side(style='thin', color='E5E7EB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    sheets = [
        ('總量管制Q&A知識庫', QA_HEADERS,
         [[r['id'], r['category'], r['keywords'], r['question'], r['answer'], r['note'], r['status']] for r in kb['qa']]),
        ('六大順位判定矩陣', kb['matrix']['headers'], kb['matrix']['rows']),
        ('學區劃分與改分發學校', kb['district']['headers'], kb['district']['rows']),
        ('115學年度重要日程表', kb['schedule']['headers'], kb['schedule']['rows']),
        ('歷年設籍門檻數據', kb['history']['headers'], kb['history']['rows']),
        ('家長在線提問紀錄簿', INQUIRY_HEADERS, []),
        ('待補充問題庫', UNANSWERED_HEADERS, []),
    ]

    first = True
    for title, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        first = False
        ws.append(headers)
        for row in rows:
            ws.append(row)

        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill, c.font, c.alignment, c.border = header_fill, header_font, align_center, border
        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 42
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font, c.border = regular_font, border
                c.alignment = align_center if len(str(c.value or '')) <= 4 else align_left
        for col in ws.columns:
            max_len = max((sum(2.2 if ord(ch) > 127 else 1.1 for ch in str(c.value or '')) for c in col), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 3, 11), 58)

    wb.save(output_path)
    print(f'Excel 建置完成：{output_path} ({os.path.getsize(output_path)} bytes)')


if __name__ == '__main__':
    with open(KB_PATH, 'r', encoding='utf-8') as f:
        kb = json.load(f)
    if os.path.isdir(GDRIVE_DIR):
        build_quota_database(kb, os.path.join(GDRIVE_DIR, FILE_NAME))
    else:
        print(f'略過雲端硬碟（找不到資料夾）：{GDRIVE_DIR}')
    build_quota_database(kb, os.path.join(BASE_DIR, FILE_NAME))
